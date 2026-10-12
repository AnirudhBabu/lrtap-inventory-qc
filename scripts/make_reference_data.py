"""Generates db/03_reference_data.sql: the countries, pollutants and sectors
of the EEA LRTAP dataset, as reference data.

Sources, in order of authority:
  docs/eea/LRTAP_Table_definition.xlsx  the EEA's code lists (sectors, pollutants, units)
  the dataset itself                    which sectors count toward the national total
                                        (column parentSectorCode), and area names
  RG2022 (docs/SOURCES.md)              which pollutants are required, from which year

The script refuses to write anything if the data and the code lists disagree,
or if the result does not match the counts CEIP publishes (127 sectors summed
into the national total, 8 memo items, 3 of them natural sources:
CEIP-M2025 p. 7).

Run:  uv run scripts/make_reference_data.py data/raw/CLRTAP_NFR19_V26_1_GF.csv
"""

import csv
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TABLE_DEFINITION = ROOT / "docs/eea/LRTAP_Table_definition.xlsx"
OUTPUT = ROOT / "db/03_reference_data.sql"

# This project's code for each pollutant name in the dataset, and the year
# from which RG2022 requires it (None: not required).
#   paras. 7(a)-(e), 38: main pollutants from 1990
#   paras. 7(f), 38:     PM2.5 and PM10 from 2000
#   paras. 7(g)-(m), 43: Cd, Pb, Hg, PAHs, PCDD/F, PCBs, HCB; from 1990 as in CEIP stage 1
#   para. 8:             BC and TSP encouraged only; other metals not listed in para. 7
POLLUTANTS = {
    # source name                    code      required   source
    "NOx":                         ("NOX",    1990, "RG2022 paras. 7(b), 38"),
    "SOx":                         ("SOX",    1990, "RG2022 paras. 7(a), 38"),
    "NMVOC":                       ("NMVOC",  1990, "RG2022 paras. 7(c), 38"),
    "NH3":                         ("NH3",    1990, "RG2022 paras. 7(d), 38"),
    "CO":                          ("CO",     1990, "RG2022 paras. 7(e), 38"),
    "PM2.5":                       ("PM2.5",  2000, "RG2022 paras. 7(f), 38"),
    "PM10":                        ("PM10",   2000, "RG2022 paras. 7(f), 38"),
    "TSP":                         ("TSP",    None, "RG2022 para. 8(b): encouraged"),
    "BC":                          ("BC",     None, "RG2022 para. 8(a): encouraged"),
    "Cd":                          ("CD",     1990, "RG2022 paras. 7, 43"),
    "Pb":                          ("PB",     1990, "RG2022 paras. 7, 43"),
    "Hg":                          ("HG",     1990, "RG2022 paras. 7, 43"),
    "As":                          ("AS",     None, "not listed in RG2022 para. 7"),
    "Cr":                          ("CR",     None, "not listed in RG2022 para. 7"),
    "Cu":                          ("CU",     None, "not listed in RG2022 para. 7"),
    "Ni":                          ("NI",     None, "not listed in RG2022 para. 7"),
    "Se":                          ("SE",     None, "not listed in RG2022 para. 7"),
    "Zn":                          ("ZN",     None, "not listed in RG2022 para. 7"),
    "Benzo(a) Pyrene":             ("BAP",    1990, "RG2022 paras. 7, 43 (PAHs)"),
    "Benzo(b) Fluoranthene":       ("BBF",    1990, "RG2022 paras. 7, 43 (PAHs)"),
    "benzo(k) Fluoranthene":       ("BKF",    1990, "RG2022 paras. 7, 43 (PAHs)"),
    "Indeno (1,2,3-cd) Pyrene":    ("IP",     1990, "RG2022 paras. 7, 43 (PAHs)"),
    "Total PAHs":                  ("PAH4",   1990, "RG2022 paras. 7, 43 (PAHs)"),
    "PCDD/PCDF (dioxins/furans)":  ("PCDD_F", 1990, "RG2022 paras. 7, 43"),
    "PCBs":                        ("PCB",    1990, "RG2022 paras. 7, 43"),
    "HCB":                         ("HCB",    1990, "RG2022 paras. 7, 43"),
}
ACTIVITY_NOTE = "activity data (fuel use), not an emission"

# Areas in the dataset that are not a reporting Party's own inventory.
OUT_OF_SCOPE_AREAS = {
    "EEA32": "EEA aggregate of its 32 member countries, not a reporting Party",
    # Checked on the 2026 edition: all 76,741 numeric EU27 cells equal the
    # sum of the 27 Member States, and its empty cells are exactly those where
    # no Member State has a number. Checking it would repeat members' findings.
    "EU27": "sum of the 27 Member States in this dataset, so its findings repeat theirs",
}

# Rows outside the national total that are not source categories.
NON_SOURCE_ROWS = {"NATIONAL TOTAL", "NATIONAL TOTAL FOR COMPLIANCE", "Adjustments"}
TOTAL_CHECKED = "NATIONAL TOTAL"


def sql(value):
    if value is None:
        return "null"
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def read_code_lists():
    wb = openpyxl.load_workbook(TABLE_DEFINITION, read_only=True)
    sectors = {
        code.strip(): name
        for code, name in wb["sectorCode - sectorName"].iter_rows(min_row=2, values_only=True)
        if code
    }
    pollutants = {
        name: unit
        for name, _code, unit in wb["pollutantName"].iter_rows(min_row=2, values_only=True)
        if name
    }
    return sectors, pollutants


def read_data(path):
    """Area names, and the set of parentSectorCode values seen per sector."""
    areas = {}
    parent = defaultdict(set)
    pollutant_names = set()
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f, delimiter="\t"):
            areas[row["countryCode"]] = row["country"]
            parent[row["sectorCode"].strip()].add(row["parentSectorCode"] or None)
            pollutant_names.add(row["pollutantName"])
    return areas, parent, pollutant_names


def sector_role(code, parents):
    if parents == {"NATIONAL TOTAL"}:
        return "SECTOR"
    if parents != {None}:
        sys.exit(f"sector {code!r} is inconsistently flagged: {parents}")
    if code == TOTAL_CHECKED:
        return "TOTAL"
    if code in NON_SOURCE_ROWS:
        return "OTHER"
    return "MEMO"


def main(data_path):
    sector_names, pollutant_units = read_code_lists()
    areas, parent, pollutant_names = read_data(data_path)

    unknown = set(parent) - set(sector_names)
    if unknown:
        sys.exit(f"sectors in the data but not in the table definition: {unknown}")
    unknown = pollutant_names - set(pollutant_units)
    if unknown:
        sys.exit(f"pollutants in the data but not in the table definition: {unknown}")

    roles = {code: sector_role(code, parent[code]) for code in sector_names if code in parent}
    counts = {r: sum(1 for v in roles.values() if v == r) for r in ("SECTOR", "MEMO", "TOTAL", "OTHER")}
    natural = [c for c, r in roles.items() if r == "MEMO" and c.startswith("11")]
    expected = {"SECTOR": 127, "MEMO": 8, "TOTAL": 1, "OTHER": 2}
    if counts != expected or len(natural) != 3:
        sys.exit(f"sector roles {counts}, natural memo items {natural}: "
                 f"expected {expected} and 3 natural sources (CEIP-M2025 p. 7)")

    lines = [
        "-- GENERATED by scripts/make_reference_data.py. Do not edit by hand.",
        "-- Sources: docs/eea/LRTAP_Table_definition.xlsx, the dataset's",
        "-- parentSectorCode column, and RG2022 (see docs/SOURCES.md).",
        f"-- Sector roles: {counts}; natural memo items: {', '.join(sorted(natural))}.",
        "",
    ]
    for code in sorted(areas):
        note = OUT_OF_SCOPE_AREAS.get(code)
        lines.append(
            "insert into ref_country (country_code, country_name, in_scope, scope_note) values ("
            f"{sql(code)}, {sql(areas[code])}, {sql('N' if note else 'Y')}, {sql(note)});")
    lines.append("")
    for name in sorted(pollutant_units, key=str.lower):
        unit = pollutant_units[name]
        if name in POLLUTANTS:
            code, required, source = POLLUTANTS[name]
            values = (code, name, unit, required, source, "Y", None)
        elif unit == "TJNCV":
            code = name.upper().replace(" ", "_")
            values = (code, name, unit, None, "EEA LRTAP table definition", "N", ACTIVITY_NOTE)
        else:
            sys.exit(f"pollutant {name!r} has no mapping in POLLUTANTS")
        lines.append(
            "insert into ref_pollutant (pollutant_code, source_name, unit, required_from_year,"
            " source_ref, in_scope, scope_note) values ("
            + ", ".join(sql(v) for v in values) + ");")
    lines.append("")
    for code in sorted(roles):
        lines.append(
            "insert into ref_sector (sector_code, sector_name, role) values ("
            f"{sql(code)}, {sql(sector_names[code])}, {sql(roles[code])});")
    lines += ["", "commit;", ""]

    OUTPUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {OUTPUT.relative_to(ROOT)}: {len(areas)} areas, "
          f"{len(pollutant_units)} pollutants, {len(roles)} sectors {counts}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
